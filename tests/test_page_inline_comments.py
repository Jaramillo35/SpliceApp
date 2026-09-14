"""Inline Comment Carryover, driven as a user would drive it — no browser.

Feeds the invented old and new reports from ``fixtures_inline_reports``
through the real upload handlers, presses the real buttons and keys, and
checks what the review gate is for: the likeliest-wrong comment is put in
front of the engineer first, deciding moves on and can be undone, there is
no bulk path through a changed row, the report waits for every row and
every comment that could not carry, and the download is the report in its
own format.
"""

from __future__ import annotations

import asyncio
import inspect
import io
import tempfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from nicegui import ui
from nicegui.testing import User
from openpyxl import load_workbook

from splice.inline import carryover as co
from tests import fixtures_inline_reports as fx


@pytest.fixture(scope="module")
def files() -> dict:
    with tempfile.TemporaryDirectory(prefix="inline_reports_") as td:
        folder = Path(td)
        old = folder / "OLD_Inline_Report.xlsx"
        new = folder / "NEW_Inline_Report.xlsx"
        old.write_bytes(fx.old_report())
        new.write_bytes(fx.new_report())
        yield {"old": old, "new": new}


def StubFile(path: Path):
    from nicegui.elements.upload_files import SmallFileUpload
    return SmallFileUpload(
        path.name,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        path.read_bytes())


def in_order(elements):
    return sorted(elements, key=lambda e: e.id)


def button(user: User, text: str) -> ui.button:
    matches = [b for b in user.find(ui.button).elements if b.text == text]
    assert len(matches) == 1, f"expected one {text!r} button, found {len(matches)}"
    return matches[0]


async def press(user: User, text: str) -> None:
    btn = button(user, text)
    assert btn.enabled, f"{text!r} is disabled"
    for listener in btn._event_listeners.values():
        if listener.type == "click":
            result = listener.handler(None)
            if inspect.isawaitable(result):
                await result
    await asyncio.sleep(0.2)


async def key(user: User, name: str) -> None:
    """A key as the browser's keyboard element reports it."""
    board = next(iter(user.find(ui.keyboard).elements))
    board._handle_key(SimpleNamespace(args={
        "action": "keydown", "repeat": False, "altKey": False, "ctrlKey": False,
        "metaKey": False, "shiftKey": False, "key": name, "code": name,
        "location": 0}))
    await asyncio.sleep(0.2)


async def wait_for(user: User, text: str, seconds: float = 30.0) -> None:
    deadline = asyncio.get_event_loop().time() + seconds
    while True:
        try:
            await user.should_see(text, retries=1)
            return
        except AssertionError:
            if asyncio.get_event_loop().time() > deadline:
                raise
            await asyncio.sleep(0.25)


async def open_matched(user: User, files: dict) -> None:
    await user.open("/inline-comments")
    await user.should_see("Match comments")
    old, new = in_order(user.find(ui.upload).elements)
    await old.handle_uploads([StubFile(files["old"])])
    await new.handle_uploads([StubFile(files["new"])])
    await wait_for(user, f"✓ {files['old'].name}")
    await wait_for(user, f"✓ {files['new'].name}")
    await press(user, "Match comments")
    await wait_for(user, "Rows to decide")


CARD_CHECK_SIZE = "X901A - Y901A · pin 4 · Q104"
CARD_SIDE_GONE = "X901A - Y901A · pin 8 · Q108 / —"
CARD_SUFFIX = "X902A - Y902A · pin 2 · Q202"
CARD_COLOUR = "X901A - Y901A · pin 3 · Q103"
CARD_FIRST_LOST = "X901A - Y901A · pin 6 · Q106"
CARD_CAVITY_12 = "X901A - Y901A · pin 12 · Q112 / —"


class TestInputs:
    async def test_match_waits_for_both_reports(self, user: User, files):
        await user.open("/inline-comments")
        await user.should_see("Match comments")
        assert not button(user, "Match comments").enabled
        await user.should_see("the OLD report")
        old, _new = in_order(user.find(ui.upload).elements)
        await old.handle_uploads([StubFile(files["old"])])
        await wait_for(user, f"✓ {files['old'].name}")
        assert not button(user, "Match comments").enabled


class TestInstructionsAndCoverage:
    async def test_the_page_says_how_it_works_before_anything_is_loaded(self, user: User):
        await user.open("/inline-comments")
        await user.should_see("How it works — and what you do")

    async def test_missing_inlines_are_shown_after_the_match(self, user: User, files):
        await open_matched(user, files)
        await user.should_see("Inlines: 2 in both reports (matched)")
        await user.should_see("Only in OLD: X903A - Y903A")
        await user.should_see("Only in NEW: X904A - Y904A")


class TestReview:
    async def test_the_result_is_counted_and_nothing_is_dropped(self, user: User, files):
        await open_matched(user, files)
        for label in ("Copied as-is", "To re-check", "Rows to decide",
                      "Comments to acknowledge"):
            await user.should_see(label)
        await user.should_see("X903A - Y903A: in the old report but not the new one")
        await user.should_see(f"{co.GROUP_LOST} · 3")

    async def test_the_likeliest_wrong_comment_is_on_the_card_first(self, user: User, files):
        await open_matched(user, files)
        await user.should_see(CARD_CHECK_SIZE)
        await user.should_see("Re-check")
        await user.should_see("It talks about Size 1, Size 2 — which changed.")

    async def test_deciding_moves_on_and_undo_brings_it_back(self, user: User, files):
        await open_matched(user, files)
        await press(user, "Leave blank")
        await wait_for(user, CARD_SIDE_GONE)
        await user.should_see("Side 2 is gone")
        await user.should_see("a decision on 6 row(s)")
        await key(user, "u")
        await wait_for(user, CARD_CHECK_SIZE)
        await wait_for(user, "a decision on 7 row(s)")

    async def test_the_keyboard_moves_and_decides(self, user: User, files):
        await open_matched(user, files)
        await key(user, "b")                  # Check size: blank, on to the side-gone row
        await wait_for(user, CARD_SIDE_GONE)
        await key(user, "ArrowDown")          # look at the next one without deciding
        await wait_for(user, CARD_CAVITY_12)  # the one-old-comment, two-rows cavity
        await user.should_see("Same cavity — 2 rows for pin 12")
        await key(user, "c")                  # copy; on to the suffix change
        await wait_for(user, CARD_SUFFIX)
        await wait_for(user, "a decision on 5 row(s)")

    async def test_there_is_no_bulk_path_through_a_changed_row(self, user: User, files):
        await open_matched(user, files)
        texts = [b.text for b in user.find(ui.button).elements]
        assert not any("changed row" in t for t in texts), texts
        await user.should_see("Copy the suggested comment on 1 sales-code row(s)")
        await user.should_see("1 sales-code row(s) whose comment talks about the "
                              "variant are left for you.")

    async def test_the_gate_waits_for_every_row_and_every_lost_comment(
            self, user: User, files):
        await open_matched(user, files)
        assert not button(user, "Write the commented report").enabled
        await user.should_see("a decision on 7 row(s)")
        await user.should_see("an acknowledgement of 3 comment(s) with no row")
        await press(user, "Copy the suggested comment on 1 sales-code row(s)")
        await wait_for(user, "a decision on 6 row(s)")
        for _ in range(6):
            await key(user, "b")
        await wait_for(user, "Needs: an acknowledgement of 3 comment(s) with no row")
        assert not button(user, "Write the commented report").enabled
        await press(user, "Mark 3 comment(s) with no row as no longer applying")
        await wait_for(user, "All 10 decided")
        assert button(user, "Write the commented report").enabled


class TestGenerate:
    async def test_the_download_is_the_report_in_its_own_format(
            self, user: User, files, monkeypatch):
        """``dress=False`` is the point: the toolkit's workbook dressing adds a
        Read Me sheet, which would make this no longer an inline report."""
        from nicegui_app import components
        handed: dict = {}

        def capture(data, filename, *, dress=True):
            handed.update(data=data, filename=filename, dress=dress)

        monkeypatch.setattr(components, "deliver", capture)
        await open_matched(user, files)
        await press(user, "Copy the suggested comment on 1 sales-code row(s)")
        await press(user, "Write new comment")            # Check size
        field = next(i for i in user.find(ui.input).elements
                     if i._props.get("label") == "New comment")
        field.set_value("size re-checked")
        await press(user, "Save comment")
        await wait_for(user, CARD_SIDE_GONE)
        await key(user, "b")          # side gone: blank
        await key(user, "c")          # cavity 12, second row: copy
        await key(user, "c")          # suffix: copy
        await key(user, "c")          # colour: copy
        await key(user, "c")          # twin A, sales code about the variant: copy
        await wait_for(user, CARD_FIRST_LOST)
        await key(user, "x")          # removed wire: no longer applies
        await key(user, "x")          # Var B, fewer rows for its cavity: no longer applies
        await key(user, "r")          # removed sheet: re-place by hand
        await wait_for(user, "All 10 decided")
        await press(user, "Write the commented report")
        await wait_for(user, "NEW_Inline_Report_commented.xlsx")
        await user.should_see("1 comment(s) to re-place by hand")
        await press(user, "NEW_Inline_Report_commented.xlsx")

        assert handed["filename"] == "NEW_Inline_Report_commented.xlsx"
        assert handed["dress"] is False
        wb = load_workbook(io.BytesIO(handed["data"]))
        assert wb.sheetnames == load_workbook(io.BytesIO(files["new"].read_bytes())).sheetnames
        ws = wb["X901A - Y901A"]
        assert [ws.cell(r, 1).value for r in range(2, 13)] == [
            "WIRE TYPE OK", "SUFFIX OK", "Varient on IP", "size re-checked",
            "Variant B", "Variant A", None, None, "OPEN ISSUE", "New view", None]
