"""Two invented inline reports, old and new, in the real report's shape.

Nothing here comes from a real programme: the inlines (X901A…), circuits
(Q101…) and sales codes (ZA1…) are made up. What is copied from a real report
is only its *shape*, because that is what the carry-over depends on:

* an ``INLINES`` index sheet whose cells hyperlink to each pair sheet;
* one sheet per inline pair, headers in row 1 — ``Comments``, the side-1
  attributes, ``Pin``, the side-2 attributes mirrored — with a 24-column
  variant carrying ``Terminal_Supplier `` twice (trailing space, no side
  digit) and a 22-column variant without it;
* lines under each table (the inline names, a 'Return To Inline List'
  hyperlink, a merged note) that are not rows;
* frozen panes and an autofilter over the table;
* numbers typed differently between exports (``1`` vs ``"1"``,
  ``0.5`` vs ``"0.50"``), which must not read as a change.

Every case the carry-over distinguishes is planted, and ``EXPECTED`` says what
each should come out as.
"""

from __future__ import annotations

import io
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SIDE1 = ["Sales Code 1", "End 1", "Twist 1", "Terminal_Supplier ", "Term Matl 1",
         "Spec 1", "Size 1", "Stripe 1", "Color 1", "Circuit 1", "Circuit Suf 1"]
SIDE2 = ["Circuit Suf 2", "Circuit 2", "Color 2", "Stripe 2", "Size 2", "Spec 2",
         "Term Matl 2", "Terminal_Supplier ", "Twist 2", "End 2", "Sales Code 2"]
HEADERS_24 = ["Comments", *SIDE1, "Pin", *SIDE2]
HEADERS_22 = [h for h in HEADERS_24 if h != "Terminal_Supplier "]

YELLOW = PatternFill(fill_type="solid", fgColor="FFFFFF00")


def wire(pin, circuit: str, codes: str = "ZA1", *, codes2: Optional[str] = None,
         color: str = "RD", color2: Optional[str] = None, stripe: str = "",
         size="0.35", size2=None, spec: str = "5ABT", suffix: str = "",
         suffix2: str = "", one_sided: bool = False, inline: str = "X901A",
         mate: str = "Y901A", supplier: str = "DE") -> Dict[str, object]:
    """One cavity's row, side 2 mirroring side 1 unless told otherwise."""
    row: Dict[Tuple[str, int], object] = {}
    side1 = {"Sales Code 1": codes, "End 1": inline, "Term Matl 1": "TIN",
             "Spec 1": spec, "Size 1": size, "Stripe 1": stripe, "Color 1": color,
             "Circuit 1": circuit, "Circuit Suf 1": suffix,
             ("Terminal_Supplier ", 1): supplier}
    side2 = {} if one_sided else {
        "Circuit Suf 2": suffix2, "Circuit 2": circuit, "Color 2": color2 or color,
        "Stripe 2": stripe, "Size 2": size if size2 is None else size2,
        "Spec 2": spec, "Term Matl 2": "TIN", "End 2": mate,
        "Sales Code 2": codes if codes2 is None else codes2,
        ("Terminal_Supplier ", 2): supplier}
    row.update({(k, 0) if isinstance(k, str) else k: v for k, v in side1.items()})
    row.update({(k, 0) if isinstance(k, str) else k: v for k, v in side2.items()})
    row[("Pin", 0)] = pin
    return row


def _write_pair(wb: Workbook, name: str, headers: List[str],
                rows: List[Tuple[str, Dict]], highlight: set = frozenset()) -> None:
    ws = wb.create_sheet(name)
    for col, header in enumerate(headers, start=1):
        ws.cell(1, col, header).font = Font(bold=True)
    pin_col = headers.index("Pin") + 1
    for r, (comment, values) in enumerate(rows, start=2):
        if comment:
            cell = ws.cell(r, 1, comment)
            if r - 2 in highlight:
                cell.fill = YELLOW
        for col, header in enumerate(headers, start=1):
            if col == 1:
                continue
            if header == "Terminal_Supplier ":
                key = (header, 1 if col < pin_col else 2)
            else:
                key = (header, 0)
            value = values.get(key)
            if value not in (None, ""):
                ws.cell(r, col, value)
    last = len(rows) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"B1:{ws.cell(1, len(headers)).column_letter}{last}"
    a, b = name.split(" - ")
    ws.cell(last + 3, 1, f"{a} Inline_{a[:-1]}")
    ws.cell(last + 4, 1, f"{b} Inline_{b[:-1]}")
    back = ws.cell(last + 6, 1, "Return To Inline List")
    back.hyperlink = "#'INLINES'!B1"
    ws.cell(last + 8, 1, "Legend: rows are cavities of the inline pair")
    ws.merge_cells(start_row=last + 8, start_column=1, end_row=last + 8, end_column=6)


def _write(sheets: List[Tuple[str, List[str], List[Tuple[str, Dict]], set]]) -> bytes:
    wb = Workbook()
    index = wb.active
    index.title = "INLINES"
    index.append(["SHEET", "INLINE", "MATING INLINE"])
    for r, (name, _h, _rows, _hl) in enumerate(sheets, start=2):
        a, b = name.split(" - ")
        cell = index.cell(r, 1, name)
        cell.hyperlink = f"#'{name}'!B1"
        index.cell(r, 2, f"{a} Inline_{a[:-1]}")
        index.cell(r, 3, f"{b} Inline_{b[:-1]}")
    index.freeze_panes = "B2"
    for name, headers, rows, highlight in sheets:
        _write_pair(wb, name, headers, rows, highlight)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def old_report() -> bytes:
    x901 = [
        ("WIRE TYPE OK", wire(1, "Q101", "ZA1", size=0.5)),            # 0 identical
        ("SUFFIX OK", wire(2, "Q102", "ZA1")),                          # 1 codes change
        ("Varient on IP", wire(3, "Q103", "ZB1")),                      # 2 colour change
        ("Check size", wire(4, "Q104", "ZA2")),                         # 3 codes + size
        ("Variant A", wire(5, "Q105", "ZC1")),                          # 4 twin A
        ("Variant B", wire(5, "Q105", "ZC2")),                          # 5 twin B
        ("Removed wire", wire(6, "Q106", "ZA1")),                       # 6 lost
        ("", wire(7, "Q107", "ZA1")),                                   # 7 no comment
        ("One sided now", wire(8, "Q108", "ZB2")),                      # 8 side dropped
        ("OPEN ISSUE", wire(9, "Q109", "ZA3")),                         # 9 highlighted
        ("Old view", wire(10, "Q110", "ZA1")),                          # 10 new has own
    ]
    x902 = [
        ("OK", wire(1, "Q201", "ZD1", inline="X902A", mate="Y902A")),
        ("Suffix", wire(2, "Q202", "ZD2", suffix="AA", inline="X902A", mate="Y902A")),
    ]
    x903 = [("Gone sheet", wire(1, "Q301", "ZE1", inline="X903A", mate="Y903A"))]
    return _write([("X901A - Y901A", HEADERS_24, x901, {9}),
                   ("X902A - Y902A", HEADERS_22, x902, set()),
                   ("X903A - Y903A", HEADERS_24, x903, set())])


def new_report() -> bytes:
    x901 = [
        ("", wire("1", "Q101", "ZA1", size="0.50")),                    # identical, typed differently
        ("", wire(2, "Q102", "ZA1/ZB2")),                               # sales codes only
        ("", wire(3, "Q103", "ZB1", color2="OG")),                      # Color 2 changed
        ("", wire(4, "Q104", "ZA3", size="0.50")),                      # codes and size
        ("", wire(5, "Q105", "ZC2")),                                   # twin B, identical
        ("", wire(5, "Q105", "ZC1&-ZC2")),                              # twin A, codes changed
        ("", wire(7, "Q107", "ZA1")),
        ("", wire(8, "Q108", "ZB2", one_sided=True)),                   # side 2 gone
        ("", wire(9, "Q109", "ZA3")),                                   # highlighted, identical
        ("New view", wire(10, "Q110", "ZA1")),                          # already commented
        ("", wire(11, "Q111", "ZA1")),                                  # brand new wire
    ]
    x902 = [
        ("", wire(1, "Q201", "ZD1", inline="X902A", mate="Y902A")),
        ("", wire(2, "Q202", "ZD2", suffix="AB", inline="X902A", mate="Y902A")),
    ]
    x904 = [("", wire(1, "Q401", "ZF1", inline="X904A", mate="Y904A"))]
    return _write([("X901A - Y901A", HEADERS_24, x901, set()),
                   ("X902A - Y902A", HEADERS_22, x902, set()),
                   ("X904A - Y904A", HEADERS_24, x904, set())])


#: what each planted case must come out as; (sheet, new row) → (kind, old comment)
EXPECTED = {
    ("X901A - Y901A", 2): ("exact", "WIRE TYPE OK"),
    ("X901A - Y901A", 3): ("sales_code", "SUFFIX OK"),
    ("X901A - Y901A", 4): ("changed", "Varient on IP"),
    ("X901A - Y901A", 5): ("changed", "Check size"),
    ("X901A - Y901A", 6): ("exact", "Variant B"),
    ("X901A - Y901A", 7): ("sales_code", "Variant A"),
    ("X901A - Y901A", 9): ("changed", "One sided now"),
    ("X901A - Y901A", 10): ("exact", "OPEN ISSUE"),
    ("X901A - Y901A", 11): ("exact", "Old view"),          # new report has its own
    ("X902A - Y902A", 2): ("exact", "OK"),
    ("X902A - Y902A", 3): ("changed", "Suffix"),
}
LOST = {("X901A - Y901A", "Removed wire"), ("X903A - Y903A", "Gone sheet")}
