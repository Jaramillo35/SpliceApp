"""The sales-code band, cell by cell, against the specification.

Every shape here was found in a real master and is reproduced on an invented
one: the anchor with a line break in it, the prose beside a code, the codes on
row 8, the package and market columns that look like codes, the variant
markers inside the band. None of the data is real.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from splice.harnesscx import band


def _sheet(wb, title, *, anchor_row=6, left="Optional Features",
           right="RELEASE/EBOM STRING", row9=None, row8=None, row7=None,
           before=None, markets=None):
    """A family sheet: package columns, the market columns, the band, then
    bookkeeping. ``markets`` puts a Markets header with those codes directly
    before the left anchor, the way every real sheet has it."""
    ws = wb.create_sheet(title)
    ws.cell(9, 3, "Made from")
    ws.cell(9, 5, "Current")
    ws.cell(anchor_row, 5, "Current Part Numbers")
    col = 6
    for text in (before or ["PC1", "PC3 AWD"]):           # look like codes, are not
        ws.cell(9, col, text)
        col += 1
    ws.cell(anchor_row, 6, "CPOS Packages")
    if markets:
        ws.cell(anchor_row, col, "Markets")
        for text in markets:
            ws.cell(9, col, text)
            col += 1
    start = col
    ws.cell(anchor_row, start, left)
    for offset, text in enumerate(row9 or []):
        if text is not None:
            ws.cell(9, start + offset, text)
    for offset, text in enumerate(row8 or []):
        if text is not None:
            ws.cell(8, start + offset, text)
    for offset, text in enumerate(row7 or []):
        if text is not None:
            ws.cell(7, start + offset, text)
    end = start + max(len(row9 or []), len(row8 or []), 1)
    ws.cell(anchor_row, end, right)
    ws.cell(9, end + 1, "should never be read")
    return ws, start, end


def _bytes(wb) -> bytes:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestParsingOneCell:
    @pytest.mark.parametrize("text, codes, ops", [
        ("AHT", ["AHT"], []),
        ("CM5/CVM", ["CM5", "CVM"], ["/"]),
        ("RS3+(CM5/CVM)", ["RS3", "CM5", "CVM"], ["+", "(", "/", ")"]),
        ("XH3=XH4", ["XH3", "XH4"], ["="]),
        ("-XZ4", ["XZ4"], ["-"]),
        ("(RDG+NAFTA)", ["RDG"], ["(", "+", ")"]),
        ("AHT (XEY)", ["AHT", "XEY"], ["(", ")"]),
        ("AHT(XEY)", ["AHT", "XEY"], ["(", ")"]),
        ("501", ["501"], []),
        ("A/B%C", [], ["/", "%"]),
    ])
    def test_codes_and_operators(self, text, codes, ops):
        assert band.parse_cell(text) == (codes, ops)

    def test_prose_beside_a_code_contributes_nothing(self):
        """``[A-Z0-9]{3}`` without word boundaries reads INC and LUD out of
        'Includes'. With them, only the codes survive."""
        codes, _ = band.parse_cell("JRK (Includes JPG)")
        assert codes == ["JRK", "JPG"]
        codes, _ = band.parse_cell("RHH/RDU (included RS4 No E12ckt)")
        assert codes == ["RHH", "RDU", "RS4"]
        codes, _ = band.parse_cell("XAK (NO XAC)")
        assert codes == ["XAK", "XAC"]

    def test_a_multi_line_cell_and_leading_spaces(self):
        assert band.parse_cell("LAS/NHZ/\nLSU/LST/BGG")[0] \
            == ["LAS", "NHZ", "LSU", "LST", "BGG"]
        assert band.parse_cell("    JWG/JRP/JRS")[0] == ["JWG", "JRP", "JRS"]

    def test_matching_is_case_sensitive_as_specified(self):
        """Masters write codes in capitals. Upper-casing the cell first turned
        the 'ckt' of 'No E12 ckt' into a sales code called CKT."""
        assert band.parse_cell("aht")[0] == []
        assert band.parse_cell("RHH/RDU (included RS4 No E12 ckt)")[0] \
            == ["RHH", "RDU", "RS4", "E12"]

    def test_duplicates_collapse_but_operators_are_counted(self):
        codes, ops = band.parse_cell("AAA/AAA/BBB")
        assert codes == ["AAA", "BBB"]
        assert ops == ["/", "/"]

    def test_the_cell_classifies_its_own_shape(self):
        cell = lambda t: band.SalesCodeCell(1, "A", t, *band.parse_cell(t))  # noqa: E731
        assert cell("AHT").is_single
        assert cell("CM5/CVM").is_or_list
        assert cell("XH3=XH4").is_equality
        assert cell("RS3+(CM5/CVM)").is_combined
        assert not cell("CM5/CVM").is_combined
        assert cell("-XZ4").is_combined

    def test_the_output_shape_is_the_one_specified(self):
        cell = band.SalesCodeCell(47, "AU", "RS3+(CM5/CVM)",
                                  *band.parse_cell("RS3+(CM5/CVM)"))
        assert cell.as_dict() == {
            "column_name": "AU",
            "raw_expression": "RS3+(CM5/CVM)",
            "sales_codes": ["RS3", "CM5", "CVM"],
            "operators": ["+", "(", "/", ")"],
        }


class TestLocatingTheBand:
    def test_only_the_columns_between_the_anchors_are_read(self):
        wb = Workbook()
        _, start, end = _sheet(wb, "DASH", row9=["ERC", "SDE", "XHZ"])
        result = band.read_master(_bytes(wb))["DASH"]
        assert (result.start_col, result.end_col) == (start, end)
        assert result.codes == ["ERC", "SDE", "XHZ"]

    def test_package_columns_are_not_codes(self):
        """PC3 and AWD are three characters and sit left of the band."""
        wb = Workbook()
        _sheet(wb, "DASH", row9=["ERC"], before=["PC1", "PC3 AWD", "PC5"])
        assert band.read_master(_bytes(wb))["DASH"].codes == ["ERC"]

    def test_market_codes_are_read_separately_and_offered_not_added(self):
        """YAA and YAC sit under 'Markets', directly before 'Optional
        Features', on every real sheet. They are not in the band — the
        engineer decides, per family, whether one becomes a column."""
        wb = Workbook()
        _sheet(wb, "DASH", row9=["ERC"], markets=["YAA", "YAC"])
        result = band.read_master(_bytes(wb))["DASH"]
        assert result.codes == ["ERC"]
        assert result.market_codes == ["YAA", "YAC"]
        assert result.left_anchor == "Optional Features"
        assert [c.as_dict()["column_name"] for c in result.markets] == ["H", "I"]

    def test_without_a_markets_header_nothing_is_offered(self):
        wb = Workbook()
        _sheet(wb, "DASH", row9=["ERC"])
        result = band.read_master(_bytes(wb))["DASH"]
        assert result.codes == ["ERC"]
        assert result.markets == []

    def test_markets_do_not_disturb_the_row_8_fallback(self):
        """Putting the markets inside the band made row 9 non-empty on the
        jumper sheet and silently lost its row-8 codes. They are separate."""
        wb = Workbook()
        _sheet(wb, "GBT JMPR", row9=[None, None], row8=["EH3", "AB7"],
               markets=["YAA", "YAC"])
        result = band.read_master(_bytes(wb))["GBT JMPR"]
        assert result.row_used == 8
        assert result.codes == ["EH3", "AB7"]
        assert result.market_codes == ["YAA", "YAC"]

    def test_release_columns_are_not_codes(self):
        wb = Workbook()
        ws, _, end = _sheet(wb, "DASH", row9=["ERC"])
        ws.cell(9, end + 2, "ABC")            # right of RELEASE/EBOM STRING
        assert band.read_master(_bytes(wb))["DASH"].codes == ["ERC"]

    def test_the_anchor_with_a_line_break_is_found(self):
        """Three sheets of the reference master spell it 'Optional\\nFeatures'.
        A substring test misses them — and they were the undetected ones."""
        wb = Workbook()
        _sheet(wb, "TRAILER TOW", left="Optional\nFeatures", row9=["SDE", "AHT(XEY)"])
        result = band.read_master(_bytes(wb))
        assert "TRAILER TOW" in result
        assert result["TRAILER TOW"].codes == ["SDE", "AHT", "XEY"]

    def test_anchor_matching_is_case_and_space_insensitive(self):
        wb = Workbook()
        _sheet(wb, "X", left="  OPTIONAL   FEATURES ", right="release/ebom  string",
               row9=["AAA"])
        assert band.read_master(_bytes(wb))["X"].codes == ["AAA"]

    def test_a_sheet_without_both_anchors_is_not_a_family(self):
        wb = Workbook()
        ws = wb.create_sheet("Summary (PC)")
        ws.cell(6, 2, "RELEASE/EBOM STRING")   # the right anchor alone
        ws.cell(9, 3, "AAA")
        _sheet(wb, "DASH", row9=["ERC"])
        assert list(band.read_master(_bytes(wb))) == ["DASH"]
        assert band.family_worksheets(_bytes(wb)) == ["DASH"]

    def test_a_template_sheet_with_anchors_but_no_codes_is_not_a_family(self):
        wb = Workbook()
        _sheet(wb, "TEMPLATE SHEET", row9=[None])
        _sheet(wb, "DASH", row9=["ERC"])
        assert "TEMPLATE SHEET" in band.read_master(_bytes(wb))
        assert band.family_worksheets(_bytes(wb)) == ["DASH"]

    def test_a_missing_release_cell_is_closed_by_the_ebom_bookkeeping(self):
        """One seat sheet of the reference master has no RELEASE/EBOM STRING
        cell; its bookkeeping starts with 'EBOM Analyst' on the row above."""
        wb = Workbook()
        ws, start, end = _sheet(wb, "SEAT 3RD ROW", right="", row9=["LEFT", "RIGHT", "JRJ", "CG3"])
        ws.cell(5, end, "EBOM Analyst")            # row 5, not row 6
        ws.cell(6, end + 3, "EBOM X-Check")
        result = band.read_master(_bytes(wb))["SEAT 3RD ROW"]
        assert result.codes == ["JRJ", "CG3"]
        assert result.end_col == end
        assert result.right_anchor == "EBOM Analyst"
        assert any("No 'RELEASE/EBOM STRING'" in n for n in result.notes)

    def test_requested_field_for_ebom_closes_the_band_one_column_early(self):
        """On most sheets 'Requested Field for EBOM' precedes the specified
        anchor by one column; row 9 there is empty, so nothing is lost and
        the specified anchor still counts as found."""
        wb = Workbook()
        ws, start, end = _sheet(wb, "DASH", row9=["ERC", "SDE"])
        ws.cell(6, end, "Requested Field for EBOM Analyst")
        ws.cell(6, end + 1, "RELEASE/EBOM STRING")
        result = band.read_master(_bytes(wb))["DASH"]
        assert result.codes == ["ERC", "SDE"]
        assert result.end_col == end
        assert result.notes == []      # the band is complete; nothing to say

    def test_anchors_in_the_wrong_order_are_not_a_band(self):
        wb = Workbook()
        ws = wb.create_sheet("Odd")
        ws.cell(6, 2, "RELEASE/EBOM STRING")
        ws.cell(6, 8, "Optional Features")
        assert band.read_master(_bytes(wb)) == {}

    def test_a_sub_header_inside_the_band_is_not_a_code(self):
        """'Standard Features' sits on row 6 mid-band; row 9 under it still
        holds codes, and the header text itself must never be one."""
        wb = Workbook()
        ws, start, _ = _sheet(wb, "RDM", row9=["SDE", "BNB", "501"])
        ws.cell(6, start + 1, "Standard Features")
        assert band.read_master(_bytes(wb))["RDM"].codes == ["SDE", "BNB", "501"]

    def test_side_markers_fall_out_and_console_variants_stay(self):
        """LEFT/RIGHT never match the token shape. CUP, CM5 and CVM do, and
        they are kept: a console variant marker is also its sales code, and
        the old detection kept them too."""
        wb = Workbook()
        _sheet(wb, "SEAT 2ND ROW", row9=["AHT", "LEFT", "RIGHT", "CUP", "CM5/CVM"])
        result = band.read_master(_bytes(wb))["SEAT 2ND ROW"]
        assert result.codes == ["AHT", "CUP", "CM5", "CVM"]

    def test_the_feature_description_comes_from_row_7(self):
        wb = Workbook()
        _sheet(wb, "DASH", row9=["AHT"], row7=["TRAILER TOW"])
        cell = band.read_master(_bytes(wb))["DASH"].cells[0]
        assert cell.feature == "TRAILER TOW"
        assert cell.column_name == "H"


class TestRowNineIsTheRuleAndRowEightIsSaid:
    def test_codes_on_row_9_are_read_from_row_9(self):
        wb = Workbook()
        _sheet(wb, "DASH", row9=["AHT"], row8=["ZZZ"])
        result = band.read_master(_bytes(wb))["DASH"]
        assert result.row_used == 9
        assert result.codes == ["AHT"]
        assert result.notes == []

    def test_an_empty_row_9_falls_back_to_row_8_and_says_so(self):
        """One jumper sheet in the reference master keeps its codes on row 8."""
        wb = Workbook()
        _sheet(wb, "GBT JMPR", row9=[None, None], row8=["EH3", "AB7"])
        result = band.read_master(_bytes(wb))["GBT JMPR"]
        assert result.row_used == 8
        assert result.codes == ["EH3", "AB7"]
        assert any("row 8" in n for n in result.notes)

    def test_a_band_with_no_codes_anywhere_is_reported_not_invented(self):
        wb = Workbook()
        _sheet(wb, "EMPTY", row9=[None])
        result = band.read_master(_bytes(wb))["EMPTY"]
        assert result.codes == []
        assert result.notes and "No sales codes" in result.notes[0]


class TestReadingAWholeMaster:
    def test_every_family_sheet_is_returned_and_nothing_else(self):
        wb = Workbook()
        wb.create_sheet("Change Log").cell(1, 1, "notes")
        _sheet(wb, "IP", row9=["AHT", "CM5/CVM"])
        _sheet(wb, "DASH", left="Optional\nFeatures", row9=["ERC"])
        result = band.read_master(_bytes(wb))
        assert sorted(result) == ["DASH", "IP"]
        assert result["IP"].as_dicts()[1]["sales_codes"] == ["CM5", "CVM"]

    def test_worksheets_can_be_restricted(self):
        wb = Workbook()
        _sheet(wb, "IP", row9=["AHT"])
        _sheet(wb, "DASH", row9=["ERC"])
        assert list(band.read_master(_bytes(wb), worksheets={"DASH"})) == ["DASH"]

    def test_an_empty_upload_is_refused(self):
        from splice.common.errors import SpliceError
        with pytest.raises(SpliceError):
            band.read_master(b"")
