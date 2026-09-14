"""Carrying inline-report comments forward, on invented reports.

``fixtures_inline_reports`` plants every case the carry-over distinguishes —
an identical row typed differently, a sales-code-only change, a colour change,
a pin carrying the same circuit twice, a side that dropped away, a removed
wire, a removed sheet, a highlighted comment, a new row that already has a
comment — and ``EXPECTED`` pins what each must become.

The output test is the one that matters most: the new report comes back with
comment cells written and every other cell, link, merge, freeze and filter
exactly as it arrived.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from splice.common.errors import SpliceError
from splice.inline import carryover as co
from tests import fixtures_inline_reports as fx


@pytest.fixture()
def old_bytes():
    return fx.old_report()


@pytest.fixture()
def new_bytes():
    return fx.new_report()


@pytest.fixture()
def result(old_bytes, new_bytes):
    return co.match(co.read_report(old_bytes), co.read_report(new_bytes))


class TestReadingAReport:
    def test_only_pair_sheets_are_tables(self, old_bytes):
        report = co.read_report(old_bytes)
        assert set(report.sheets) == {"X901A - Y901A", "X902A - Y902A", "X903A - Y903A"}
        assert report.skipped == ["INLINES"]

    def test_the_lines_under_a_table_are_not_rows(self, old_bytes):
        sheet = co.read_report(old_bytes).sheets["X901A - Y901A"]
        assert len(sheet.rows) == 14
        assert all(r.pin for r in sheet.rows)

    def test_the_unsided_duplicate_header_takes_its_side_from_pin(self, old_bytes):
        keys = [c.key for c in co.read_report(old_bytes).sheets["X901A - Y901A"].columns]
        assert "Terminal_Supplier|1" in keys and "Terminal_Supplier|2" in keys
        assert keys.index("Terminal_Supplier|1") < keys.index("Pin") \
            < keys.index("Terminal_Supplier|2")

    def test_a_sheet_without_that_column_still_reads(self, old_bytes):
        sheet = co.read_report(old_bytes).sheets["X902A - Y902A"]
        assert "Terminal_Supplier|1" not in [c.key for c in sheet.columns]
        assert len(sheet.rows) == 2

    @pytest.mark.parametrize("a, b", [(1, "1"), (0.5, "0.50"), ("501", 501),
                                      (" GWJ ", "GWJ"), (None, "")])
    def test_the_same_value_typed_differently_is_the_same(self, a, b):
        assert co.normalize(a) == co.normalize(b)

    def test_a_workbook_with_no_pair_sheet_is_refused_in_words(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(["Something", "Else"])
        buffer = io.BytesIO()
        wb.save(buffer)
        with pytest.raises(SpliceError, match="inline report"):
            co.read_report(buffer.getvalue(), "notes.xlsx")


class TestEveryPlantedCase:
    def test_each_row_is_classified_as_planted(self, result):
        got = {(p.sheet, p.new_row): (p.kind, p.old_comment) for p in result.proposals}
        assert got == fx.EXPECTED

    def test_nothing_is_dropped_silently(self, result):
        assert {(lost.sheet, lost.comment) for lost in result.lost} == fx.LOST
        assert any("X903A - Y903A" in note for note in result.notes)

    def test_identical_rows_are_decided_without_asking(self, result):
        for p in result.of_kind(co.EXACT):
            assert p.decided
            assert p.decision == (co.KEEP if p.existing else co.COPY)

    def test_a_sales_code_change_suggests_the_old_comment_and_waits(self, result):
        for p in result.of_kind(co.SALES_CODE):
            assert p.suggested == co.COPY and not p.decided
            assert all(co.is_sales_code(d.key) for d in p.diffs)

    def test_any_other_change_suggests_nothing(self, result):
        """A 'WIRE TYPE OK' after the wire changed is exactly the comment that
        must not carry on its own."""
        for p in result.of_kind(co.CHANGED):
            assert p.suggested is None and not p.decided
            assert any(not co.is_sales_code(d.key) for d in p.diffs)

    def test_what_changed_is_named(self, result):
        colour = result.get("X901A - Y901A!4")
        assert colour.changed == "Color 2: RD → OG"
        one_sided = result.get("X901A - Y901A!9")
        assert "Circuit 2: Q108 → —" in one_sided.changed
        suffix = result.get("X902A - Y902A!3")
        assert suffix.changed == "Circuit Suf 1: AA → AB"

    def test_twins_on_one_pin_each_keep_their_own_comment(self, result):
        """Pin 5 carries Q105 twice. The identical twin is paired first, so the
        other twin cannot steal its comment."""
        assert result.get("X901A - Y901A!6").old_comment == "Variant B"
        assert result.get("X901A - Y901A!7").old_comment == "Variant A"

    def test_a_row_that_already_has_a_comment_keeps_it_by_default(self, result):
        p = result.get("X901A - Y901A!11")
        assert p.existing == "New view" and p.decision == co.KEEP

    def test_a_brand_new_row_gets_nothing(self, result):
        assert all(p.new_row != 12 for p in result.proposals if p.sheet == "X901A - Y901A")

    def test_counts(self, result):
        n = result.counts()
        assert (n[co.EXACT], n[co.SALES_CODE], n[co.CHANGED]) == (7, 2, 5)
        assert n["undecided"] == 7 and n["lost"] == 3 and n["kept"] == 1


class TestTheReviewGate:
    def test_the_report_is_not_written_while_rows_are_undecided(
            self, result, old_bytes, new_bytes):
        with pytest.raises(co.UndecidedRows) as caught:
            co.apply(old_bytes, new_bytes, result)
        assert caught.value.count == 7

    def test_accepting_suggestions_touches_only_rows_that_have_one(self, result):
        assert result.accept_suggestions() == 2
        assert all(p.decided for p in result.of_kind(co.SALES_CODE))
        assert not any(p.decided for p in result.of_kind(co.CHANGED))

    def test_a_new_comment_needs_text(self, result):
        with pytest.raises(ValueError):
            result.decide("X901A - Y901A!4", co.NEW, "   ")
        assert result.decide("X901A - Y901A!4", co.NEW, " colour check ").text == "colour check"

    def test_keep_is_only_for_a_row_with_its_own_comment(self, result):
        with pytest.raises(ValueError):
            result.decide("X901A - Y901A!4", co.KEEP)

    def test_an_identical_row_cannot_be_sent_back_to_undecided(self, result):
        assert result.undo("X901A - Y901A!2").decided
        result.accept_suggestions()
        assert not result.undo("X901A - Y901A!3").decided

    def test_bulk_decisions_cannot_write_new_text(self, result):
        with pytest.raises(ValueError):
            result.decide_all(co.CHANGED, co.NEW)
        assert result.decide_all(co.CHANGED, co.BLANK) == 5


def _decide_everything_but_lost(result):
    result.accept_suggestions()
    result.decide("X901A - Y901A!4", co.COPY)
    result.decide("X901A - Y901A!5", co.NEW, "size re-checked")
    result.decide("X901A - Y901A!9", co.BLANK)
    result.decide("X902A - Y902A!3", co.COPY)
    result.decide_cavity("X901A - Y901A!13", co.COPY)     # both rows of cavity 12


def _decide_everything(result):
    _decide_everything_but_lost(result)
    result.acknowledge_all(co.OBSOLETE)


class TestTheOutput:
    def test_comments_are_written_where_decided(self, result, old_bytes, new_bytes):
        _decide_everything(result)
        ws = load_workbook(io.BytesIO(co.apply(old_bytes, new_bytes, result)))["X901A - Y901A"]
        assert [ws.cell(r, 1).value for r in range(2, 16)] == [
            "WIRE TYPE OK", "SUFFIX OK", "Varient on IP", "size re-checked",
            "Variant B", "Variant A", None, None, "OPEN ISSUE", "New view", None,
            "Splice both", "Splice both", "Var A"]

    def test_everything_else_is_the_new_report_unchanged(
            self, result, old_bytes, new_bytes):
        _decide_everything(result)
        before = load_workbook(io.BytesIO(new_bytes))
        after = load_workbook(io.BytesIO(co.apply(old_bytes, new_bytes, result)))
        assert after.sheetnames == before.sheetnames, "no sheet added or removed"
        written = {(p.sheet, p.new_row) for p in result.proposals}
        for name in before.sheetnames:
            a, b = before[name], after[name]
            assert (a.freeze_panes, a.auto_filter.ref) == (b.freeze_panes, b.auto_filter.ref)
            assert sorted(map(str, a.merged_cells.ranges)) == sorted(map(str, b.merged_cells.ranges))
            assert (a.max_row, a.max_column) == (b.max_row, b.max_column)
            for row in a.iter_rows():
                for cell in row:
                    if cell.column == 1 and (name, cell.row) in written:
                        continue
                    other = b[cell.coordinate]
                    assert cell.value == other.value, cell.coordinate
                    assert bool(cell.hyperlink) == bool(other.hyperlink), cell.coordinate
                    assert cell.font.b == other.font.b, cell.coordinate

    def test_a_copied_comment_keeps_its_highlight(self, result, old_bytes, new_bytes):
        _decide_everything(result)
        ws = load_workbook(io.BytesIO(co.apply(old_bytes, new_bytes, result)))["X901A - Y901A"]
        assert ws.cell(10, 1).value == "OPEN ISSUE"
        assert ws.cell(10, 1).fill.fgColor.rgb == "FFFFFF00"
        assert ws.cell(2, 1).fill.fill_type is None

    def test_carrying_a_report_onto_itself_copies_every_comment(self, old_bytes):
        report = co.read_report(old_bytes)
        blank = io.BytesIO()
        wb = load_workbook(io.BytesIO(old_bytes))
        for name in report.sheets:
            for row in report.sheets[name].rows:
                wb[name].cell(row.row, 1).value = None
        wb.save(blank)
        result = co.match(report, co.read_report(blank.getvalue()))
        assert result.counts()[co.EXACT] == report.commented
        assert not result.undecided and not result.lost
        out = load_workbook(io.BytesIO(co.apply(old_bytes, blank.getvalue(), result)))
        for name, sheet in report.sheets.items():
            for row in sheet.rows:
                assert (out[name].cell(row.row, 1).value or "") == row.comment

    @pytest.mark.parametrize("name, expected", [
        ("IP_Inline_Report.xlsx", "IP_Inline_Report_commented.xlsx"),
        ("5.- IP_11661_Inline_Report.xlsx", "5.- IP_11661_Inline_Report_commented.xlsx"),
        ("report", "report_commented.xlsx"),
    ])
    def test_output_name(self, name, expected):
        assert co.output_name(name) == expected


class TestTriage:
    """A better gate spends the engineer's attention where a comment is most
    likely to have stopped being true."""

    @pytest.mark.parametrize("comment, key, old, new, expected", [
        ("Check size", "Size|1", "0.35", "0.5", True),
        ("WIRE TYPE OK", "Spec|2", "5ABT", "5ABX", True),
        ("SUFFIX OK", "Sales Code|1", "ZA1", "ZA1/ZB2", False),
        ("Varient on IP", "Color|2", "RD", "OG", False),
        ("RD per drawing", "Color|2", "RD", "OG", True),          # by value
        ("Oversize boot", "Size|1", "0.35", "0.5", False),        # a word, not a fragment
        ("Variant A", "Sales Code|2", "ZC1", "ZC1&-ZC2", True),
        ("", "Size|1", "0.35", "0.5", False),
    ])
    def test_a_comment_mentions_what_changed(self, comment, key, old, new, expected):
        assert bool(co.mentions(comment, [co.Diff(key, old, new)])) is expected

    def test_each_planted_row_lands_in_its_group(self, result):
        groups = {p.id: p.group for p in result.proposals}
        assert groups["X901A - Y901A!5"] == co.GROUP_RECHECK        # 'Check size', size changed
        assert groups["X901A - Y901A!9"] == co.GROUP_RECHECK        # side 2 is gone
        assert groups["X902A - Y902A!3"] == co.GROUP_RECHECK        # 'Suffix', suffix changed
        assert groups["X901A - Y901A!4"] == co.GROUP_CHANGED        # colour moved, comment is about the variant
        assert groups["X901A - Y901A!7"] == co.GROUP_CODES_RECHECK  # 'Variant A', codes moved
        assert groups["X901A - Y901A!3"] == co.GROUP_CODES
        assert groups["X901A - Y901A!11"] == co.GROUP_KEPT
        assert groups["X901A - Y901A!2"] == co.GROUP_IDENTICAL

    def test_a_side_that_is_gone_is_named(self, result):
        assert result.get("X901A - Y901A!9").sides_gone == ["2"]
        assert result.get("X901A - Y901A!5").sides_gone == []

    def test_the_queue_puts_the_riskiest_first_and_hides_what_needs_no_one(self, result):
        ids = [x.id for x in result.queue()]
        assert ids == ["X901A - Y901A!5", "X901A - Y901A!9", "X901A - Y901A!14",
                       "X902A - Y902A!3",
                       "X901A - Y901A!4",
                       "X901A - Y901A!7",
                       "X901A - Y901A!3",
                       "X901A - Y901A!8!lost", "X901A - Y901A!15!lost",
                       "X903A - Y903A!2!lost"]
        settled = {x.id for x in result.queue(include_settled=True)} - set(ids)
        assert settled == {"X901A - Y901A!2", "X901A - Y901A!6", "X901A - Y901A!10",
                           "X901A - Y901A!11", "X901A - Y901A!13", "X901A - Y901A!15",
                           "X902A - Y902A!2"}

    def test_deciding_a_row_does_not_move_the_list(self, result):
        before = [x.id for x in result.queue()]
        result.decide("X901A - Y901A!5", co.BLANK)
        result.acknowledge("X903A - Y903A!2!lost", co.OBSOLETE)
        assert [x.id for x in result.queue()] == before


class TestLostCommentsAreAcknowledged:
    """'Nothing is dropped silently' holds only if the gate makes someone look
    at the comments that could not carry."""

    def test_the_report_waits_for_every_lost_comment_to_be_seen(
            self, result, old_bytes, new_bytes):
        _decide_everything_but_lost(result)
        assert not result.undecided
        with pytest.raises(co.UndecidedRows) as caught:
            co.apply(old_bytes, new_bytes, result)
        assert (caught.value.count, caught.value.lost) == (0, 3)
        assert "did not carry" in str(caught.value)

    def test_acknowledging_clears_the_gate_without_overwriting(
            self, result, old_bytes, new_bytes):
        _decide_everything_but_lost(result)
        result.acknowledge("X901A - Y901A!8!lost", co.REPLACE)
        assert result.blocking == 2
        assert result.acknowledge_all(co.OBSOLETE) == 2
        assert result.blocking == 0
        co.apply(old_bytes, new_bytes, result)
        assert result.get_lost("X901A - Y901A!8!lost").ack == co.REPLACE

    def test_an_acknowledgement_can_be_taken_back(self, result):
        result.acknowledge("X903A - Y903A!2!lost", co.OBSOLETE)
        result.unacknowledge("X903A - Y903A!2!lost")
        assert len(result.unacknowledged) == 3

    def test_only_known_acknowledgements(self, result):
        with pytest.raises(ValueError):
            result.acknowledge("X903A - Y903A!2!lost", "whatever")
        with pytest.raises(ValueError):
            result.acknowledge_all("whatever")



class TestInlineCoverage:
    """Matching runs on the inlines present in both reports; the ones in only
    one are named, never silently skipped."""

    def test_shared_and_missing_inlines_are_named(self, result):
        cov = result.coverage
        assert cov.in_both == ["X901A - Y901A", "X902A - Y902A"]
        assert cov.old_only == ["X903A - Y903A"]
        assert cov.new_only == ["X904A - Y904A"]
        assert not cov.complete and cov.missing == 2

    def test_only_shared_inlines_produce_proposals(self, result):
        assert {p.sheet for p in result.proposals} <= set(result.coverage.in_both)

    def test_identical_reports_are_complete(self, old_bytes):
        report = co.read_report(old_bytes)
        cov = co.match(report, report).coverage
        assert cov.complete and cov.in_both == list(report.sheets)



class TestOneCavityManyRows:
    """A cavity's variants are several rows, one side blank on the extras.
    One old row → many new rows: the comment is offered on each, one per
    row. Many old rows → one new row: the extra comments are offered as
    alternatives and still listed as not carried."""

    def test_one_old_comment_reaches_every_new_row_of_the_cavity(self, result):
        a, b = result.get("X901A - Y901A!13"), result.get("X901A - Y901A!14")
        assert a.old_comment == b.old_comment == "Splice both"
        assert a.old_row == b.old_row
        assert a.kind == co.EXACT and a.decided
        assert b.kind == co.CHANGED and not b.decided and b.shared and a.shared
        assert "Side 2 is gone" not in b.changed and "Circuit 2: Q112 → —" in b.changed

    def test_the_cavity_is_seen_together(self, result):
        mates = result.cavity_mates("X901A - Y901A!14")
        assert [m.id for m in mates] == ["X901A - Y901A!13"]
        assert result.cavity_mates("X901A - Y901A!2") == []

    def test_many_old_rows_onto_one_new_row_keeps_one_comment_per_row(self, result):
        p = result.get("X901A - Y901A!15")
        assert p.old_comment == "Var A" and p.alternatives == ["Var B"]
        lost = next(x for x in result.lost if x.comment == "Var B")
        assert "fewer rows for this cavity" in lost.reason
        assert [m.id for m in result.cavity_mates(p.id)] == [lost.id]

    def test_a_cavity_can_be_decided_as_a_whole(self, result):
        assert result.decide_cavity("X901A - Y901A!14", co.COPY) == 1
        assert result.get("X901A - Y901A!14").result == "Splice both"
        with pytest.raises(ValueError):
            result.decide_cavity("X901A - Y901A!14", co.NEW)

    def test_a_single_row_cavity_is_untouched_by_the_rule(self, result):
        """Pin 5 has two rows on both sides — the one-to-one pass still pairs
        each twin with its own comment; nothing is shared."""
        assert not result.get("X901A - Y901A!6").shared
        assert not result.get("X901A - Y901A!7").shared
