"""Set the Terminal Material of circuits in DEF Editor from an Excel list.

The engineer navigates DEF Editor by hand to Edit Harness → Circuits, where
every circuit of the harness is a row with a ``Term Matl`` cell that opens a
list (BERYLLIUM, GOLD, SILVER, SILVER+NICKEL, SILVER+TIN, TIN). They then
upload an Excel file with three columns — ``CNUM``, ``Circuit Name``,
``Terminal`` — and this sets each matching row's Term Matl.

This is the one thing in the kit that *changes* DEF Editor, so it is built as
preview, then apply, then report:

**Preview** matches every Excel row to the grid without touching it — CNUM
against the ``Connector No`` column, Circuit Name against ``Circuit``, both
exactly after trimming and case — and says per row what would happen:
*change* (with the current value), *already* that value, *not found*,
*several rows* (all of them will be set, and the count is shown), *no value*
(Terminal empty: left as it is), or *unknown value*.

**Apply** sets only the rows the preview said *change*, one at a time,
reading each cell back after setting it; a cell that reads back different is
reported as *failed*, not assumed.

Values are matched exactly, never guessed. The Excel may say Silver, Gold or
Tin (any case); those become SILVER, GOLD, TIN. Anything else — including
DEF Editor's own SILVER+NICKEL, which the source list does not use — is
reported and left alone, because a wrong plating written into a released
harness is worse than a row an engineer has to do by hand.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Callable, List, Optional, Sequence

from defauto import ids
from defauto.backend import AutomationError, Backend, Grid

#: what the incoming Excel may say, and what DEF Editor calls it
INPUT_VALUES = {"SILVER": "SILVER", "GOLD": "GOLD", "TIN": "TIN"}

#: the Excel's columns, matched case-insensitively and trimmed
XLSX_COLUMNS = ("CNUM", "Circuit Name", "Terminal")

CHANGE = "change"
ALREADY = "already"
NOT_FOUND = "not found"
NO_VALUE = "no value"
UNKNOWN = "unknown value"
APPLIED = "applied"
FAILED = "failed"

STATUS_LABEL = {
    CHANGE: "Will change", ALREADY: "Already that value",
    NOT_FOUND: "Not in the grid", NO_VALUE: "Terminal empty — left as is",
    UNKNOWN: "Unknown value — left as is", APPLIED: "Applied",
    FAILED: "Failed — read back different",
}


def normalise_terminal(value) -> Optional[str]:
    """``' silver '`` → ``SILVER``; empty → ``""``; anything else → None."""
    text = "" if value is None else str(value).strip().upper()
    if not text:
        return ""
    return INPUT_VALUES.get(text)


@dataclass
class Update:
    """One row of the Excel."""

    line: int                   # the Excel row number, for the report
    cnum: str
    circuit: str
    terminal: str               # as written


@dataclass
class Planned:
    update: Update
    status: str
    target: str = ""            # the DEF Editor value to set
    current: str = ""           # what the cell holds now
    rows: List[int] = field(default_factory=list)   # grid row indexes (0-based)
    detail: str = ""

    @property
    def will_change(self) -> bool:
        return self.status == CHANGE

    def as_row(self) -> dict:
        return {"Excel row": self.update.line, "CNUM": self.update.cnum,
                "Circuit": self.update.circuit, "Terminal": self.update.terminal,
                "Current": self.current, "Set to": self.target,
                "Rows": len(self.rows), "Status": STATUS_LABEL.get(self.status, self.status),
                "Detail": self.detail}


# ------------------------------------------------------------------ reading
def read_updates(data: bytes, name: str = "") -> List[Update]:
    """The Excel's rows. Column order does not matter; the three names do."""
    from openpyxl import load_workbook  # noqa: PLC0415 - only this reads Excel

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - said in the user's words
        raise AutomationError(f"{name or 'The file'} could not be opened as an "
                              f"Excel workbook: {exc}") from exc
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise AutomationError(f"{name or 'The file'} is empty")
    labels = ["" if h is None else str(h).strip().lower() for h in header]
    where = {}
    for wanted in XLSX_COLUMNS:
        try:
            where[wanted] = labels.index(wanted.lower())
        except ValueError:
            raise AutomationError(
                f"{name or 'The file'} has no {wanted!r} column. It needs "
                f"{', '.join(XLSX_COLUMNS)} in its first row; it has "
                f"{', '.join(l for l in labels if l) or 'nothing'}.") from None
    out: List[Update] = []
    for line, values in enumerate(rows, start=2):
        cell = lambda key: (  # noqa: E731
            "" if values[where[key]] is None else str(values[where[key]]).strip())
        if not any(v not in (None, "") for v in values):
            continue
        out.append(Update(line, cell("CNUM"), cell("Circuit Name"), cell("Terminal")))
    if not out:
        raise AutomationError(f"{name or 'The file'} has a header and no rows")
    return out


# ---------------------------------------------------------------- planning
def _column(grid: Grid, wanted: str) -> int:
    """Index of a grid column by its header, tolerant of case and spacing."""
    key = wanted.strip().lower().replace(" ", "")
    for i, header in enumerate(grid.headers):
        if header.strip().lower().replace(" ", "") == key:
            return i
    raise AutomationError(
        f"The circuits grid has no {wanted!r} column — its headers are: "
        f"{', '.join(grid.headers) or 'none read'}. Is DEF Editor on the "
        "Edit Harness → Circuits page?")


def plan(backend: Backend, updates: Sequence[Update]) -> List[Planned]:
    """Match every Excel row to the grid, touching nothing."""
    grid = backend.grid(ids.GRID_CIRCUITS)
    cnum_i = _column(grid, ids.TERM_COLUMNS["cnum"])
    ckt_i = _column(grid, ids.TERM_COLUMNS["circuit"])
    term_i = _column(grid, ids.TERM_COLUMNS["terminal"])

    def key(cnum: str, circuit: str) -> tuple:
        return cnum.strip().upper(), circuit.strip().upper()

    index: dict = {}
    for r, row in enumerate(grid.rows):
        if max(cnum_i, ckt_i, term_i) < len(row):
            index.setdefault(key(row[cnum_i], row[ckt_i]), []).append(r)

    out: List[Planned] = []
    for u in updates:
        target = normalise_terminal(u.terminal)
        rows = index.get(key(u.cnum, u.circuit), [])
        current = sorted({grid.rows[r][term_i].strip() for r in rows}) if rows else []
        shown = " / ".join(v or "(empty)" for v in current)
        if target is None:
            out.append(Planned(u, UNKNOWN, "", shown, rows,
                               f"{u.terminal!r} is not Silver, Gold or Tin"))
        elif not rows:
            out.append(Planned(u, NOT_FOUND, target, "", [],
                               "no grid row has this CNUM and circuit"))
        elif target == "":
            out.append(Planned(u, NO_VALUE, "", shown, rows))
        elif all(v.upper() == target for v in current):
            out.append(Planned(u, ALREADY, target, shown, rows))
        else:
            detail = f"{len(rows)} rows share this CNUM and circuit — all will be set" \
                if len(rows) > 1 else ""
            out.append(Planned(u, CHANGE, target, shown, rows, detail))
    return out


def summary(planned: Sequence[Planned]) -> dict:
    out = {s: 0 for s in STATUS_LABEL}
    for p in planned:
        out[p.status] = out.get(p.status, 0) + 1
    return out


# ----------------------------------------------------------------- applying
def apply(backend: Backend, planned: Sequence[Planned],
          on_row: Optional[Callable[[Planned], None]] = None) -> List[Planned]:
    """Set every row the preview said *change*, reading each back.

    Rows the preview did not mark as a change are never touched — a preview
    is a promise about what apply will do, and apply keeps it.
    """
    term = ids.TERM_COLUMNS["terminal"]
    for p in planned:
        if not p.will_change:
            continue
        failures = []
        for r in p.rows:
            try:
                backend.set_cell(ids.GRID_CIRCUITS, r, term, p.target)
            except AutomationError as exc:
                failures.append(f"row {r + 1}: {exc}")
                continue
        after = backend.grid(ids.GRID_CIRCUITS)
        term_i = _column(after, term)
        read_back = [after.rows[r][term_i].strip() if r < len(after.rows) else "?"
                     for r in p.rows]
        wrong = [v for v in read_back if v.upper() != p.target]
        if failures or wrong:
            p.status = FAILED
            p.detail = "; ".join(failures + [f"reads back {v or '(empty)'!r}" for v in wrong])
        else:
            p.status = APPLIED
            p.current = p.target
        if on_row is not None:
            on_row(p)
    return list(planned)


def results_csv(planned: Sequence[Planned]) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(planned[0].as_row()) if planned else
                            ["Excel row", "CNUM", "Circuit", "Terminal", "Current",
                             "Set to", "Rows", "Status", "Detail"],
                            lineterminator="\n")
    writer.writeheader()
    for p in planned:
        writer.writerow(p.as_row())
    return buf.getvalue()
