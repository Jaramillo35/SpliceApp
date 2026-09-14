"""Carry engineer comments from a past inline report into the new one.

An inline report is one workbook per release: an ``INLINES`` index sheet and
one sheet per inline pair (``X301A - Y301A``), each a table with a
``Comments`` column, the side-1 attributes, ``Pin``, and the side-2
attributes mirrored. Engineers comment a report row by row; the next release
arrives as a fresh report with every comment empty. This module moves the
comments across and makes the engineer decide only what actually needs
deciding.

How a comment finds its row, strictest first:

**Identical row** — every attribute on both sides is the same. The comment is
still true, so it is copied without asking.

**Sales code changed** — same pin, same circuit, only the ``Sales Code``
columns differ. The wire did not change, so the old comment is *suggested*;
the engineer confirms it or writes a new one.

**Row changed** — same pin and circuit on at least one side, but a colour,
size, spec, suffix or side has changed. A comment like "WIRE TYPE OK" may no
longer be true, so nothing is suggested: the engineer decides.

Matching is one-to-one. A pin can carry the same circuit twice under
different sales codes, each with its own comment; identical rows are paired
first, then the remaining rows by fewest differing columns, so a comment never
lands on two rows or on the wrong twin.

Nothing is dropped silently. A commented row with no counterpart in the new
report — a removed wire, a removed inline sheet — is listed as a comment that
did not carry. A new-report row that already has a different comment keeps
its own unless the engineer says otherwise.

The output is the NEW workbook with only comment cells written: the format,
the hyperlinks, the frozen panes and the filters are whatever the new report
already had. A copied comment brings its cell fill and font, because a
highlighted comment is highlighted for a reason.
"""

from __future__ import annotations

import io
import re
from copy import copy
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from splice.common.errors import SpliceError

# ---------------------------------------------------------------- vocabulary
EXACT = "exact"
SALES_CODE = "sales_code"
CHANGED = "changed"

COPY = "copy"
NEW = "new"
BLANK = "blank"
KEEP = "keep"

KIND_LABEL = {EXACT: "Identical row", SALES_CODE: "Sales code changed",
              CHANGED: "Row changed"}
DECISION_LABEL = {COPY: "Copy old comment", NEW: "Write new comment",
                  BLANK: "Leave blank", KEEP: "Keep new report's comment"}

_SIDED = re.compile(r"^(.*\S)\s+([12])$")
_NUMBER = re.compile(r"^-?\d+(\.\d+)?$")


class UndecidedRows(SpliceError):
    """The review gate: a report is not written while rows await a decision."""

    def __init__(self, count: int) -> None:
        super().__init__(f"{count} row(s) still need a decision before the "
                         "report can be written")
        self.count = count


def normalize(value) -> str:
    """A cell as comparable text.

    Excel stores the same pin as ``1`` in one export and ``"1"`` in the next,
    and a size as ``0.5`` or ``"0.50"``. Those are the same attribute, and a
    row must not read as changed because of how a number was typed.
    """
    if value is None:
        return ""
    text = str(value).strip() if not isinstance(value, (int, float)) else repr(value)
    if isinstance(value, bool):
        return text
    if isinstance(value, (int, float)) or _NUMBER.match(text):
        try:
            return format(Decimal(text).normalize(), "f")
        except InvalidOperation:
            return text
    return text


def label_of(key: str) -> str:
    """``Sales Code|1`` → ``Sales Code 1``; ``Pin`` stays ``Pin``."""
    if "|" not in key:
        return key
    name, side = key.split("|", 1)
    return f"{name} {side.split('#')[0]}"


def is_sales_code(key: str) -> bool:
    return key.lower().startswith("sales code|")


# ------------------------------------------------------------------ reading
@dataclass(frozen=True)
class Column:
    index: int      # 1-based
    key: str        # 'Circuit|1', 'Pin', 'Terminal_Supplier|2'
    label: str      # the header as written, stripped


def columns_of(headers) -> Optional[Tuple[int, List[Column]]]:
    """``(comment column, attribute columns)`` of a pair sheet, or None.

    Columns are keyed by name and side rather than position, because the
    sheets of one report do not share a column set (some carry
    ``Terminal_Supplier``, some do not), and that header appears twice with
    no side digit — its side is whichever side of ``Pin`` it sits on.
    """
    labels = ["" if h is None else str(h).strip() for h in headers]
    lower = [text.lower() for text in labels]
    if "pin" not in lower or "comments" not in lower:
        return None
    pin = lower.index("pin")
    comment = lower.index("comments")
    columns: List[Column] = []
    seen = set()
    for i, text in enumerate(labels):
        if i == comment or not text:
            continue
        if i == pin:
            key = "Pin"
        else:
            sided = _SIDED.match(text)
            key = (f"{sided.group(1)}|{sided.group(2)}" if sided
                   else f"{text}|{1 if i < pin else 2}")
        if key in seen:                       # defensive; not seen in real reports
            key = f"{key}#{i + 1}"
        seen.add(key)
        columns.append(Column(i + 1, key, text))
    return comment + 1, columns


@dataclass
class Row:
    sheet: str
    row: int
    values: Dict[str, str]
    comment: str = ""

    @property
    def pin(self) -> str:
        return self.values.get("Pin", "")

    @property
    def circuit1(self) -> str:
        return self.values.get("Circuit|1", "")

    @property
    def circuit2(self) -> str:
        return self.values.get("Circuit|2", "")


@dataclass
class Sheet:
    name: str
    comment_col: int
    columns: List[Column]
    rows: List[Row] = field(default_factory=list)


@dataclass
class Report:
    sheets: Dict[str, Sheet] = field(default_factory=dict)
    #: sheets that are not pair tables (the INLINES index, anything else)
    skipped: List[str] = field(default_factory=list)

    @property
    def rows(self) -> int:
        return sum(len(s.rows) for s in self.sheets.values())

    @property
    def commented(self) -> int:
        return sum(1 for s in self.sheets.values() for r in s.rows if r.comment)


def read_report(data: bytes, name: str = "") -> Report:
    """Every inline-pair table in the workbook.

    A row is part of the table when any attribute column has a value; the
    lines under each table — the inline names, 'Return To Inline List' —
    carry text only in the Comments column and are not rows.
    """
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 — said in the user's words
        raise SpliceError(f"{name or 'The file'} could not be opened as an "
                          f"Excel workbook: {exc}") from exc
    report = Report()
    for ws in wb.worksheets:
        if ws.max_row < 1:
            report.skipped.append(ws.title)
            continue
        shape = columns_of([c.value for c in ws[1]])
        if shape is None:
            report.skipped.append(ws.title)
            continue
        comment_col, columns = shape
        sheet = Sheet(ws.title, comment_col, columns)
        for r in range(2, ws.max_row + 1):
            values = {c.key: normalize(ws.cell(r, c.index).value) for c in columns}
            if not any(values.values()):
                continue
            raw = ws.cell(r, comment_col).value
            sheet.rows.append(Row(ws.title, r, values,
                                  "" if raw is None else str(raw).strip()))
        report.sheets[ws.title] = sheet
    if not report.sheets:
        raise SpliceError(f"{name or 'The file'} has no inline-pair sheets — no "
                          "sheet has both a 'Comments' and a 'Pin' column in row 1. "
                          "Is this an inline report?")
    return report


# ----------------------------------------------------------------- matching
@dataclass
class Diff:
    key: str
    old: str
    new: str

    @property
    def label(self) -> str:
        return label_of(self.key)

    def __str__(self) -> str:
        return f"{self.label}: {self.old or '—'} → {self.new or '—'}"


@dataclass
class Proposal:
    """One new-report row that a past comment could apply to."""

    sheet: str
    new_row: int
    old_row: int
    new_col: int
    old_col: int
    kind: str
    diffs: List[Diff]
    old_comment: str
    existing: str
    pin: str
    circuit1: str
    circuit2: str
    suggested: Optional[str]
    decision: Optional[str]
    text: str = ""
    #: comments of other old rows that matched this one equally well
    alternatives: List[str] = field(default_factory=list)

    @property
    def id(self) -> str:
        return f"{self.sheet}!{self.new_row}"

    @property
    def decided(self) -> bool:
        return self.decision is not None

    @property
    def changed(self) -> str:
        return "; ".join(str(d) for d in self.diffs)

    @property
    def result(self) -> str:
        """What the cell will hold once written."""
        return {COPY: self.old_comment, NEW: self.text, BLANK: "",
                KEEP: self.existing}.get(self.decision or "", "")


@dataclass
class Lost:
    """A comment in the old report that no new row can take."""

    sheet: str
    row: int
    comment: str
    pin: str
    circuit1: str
    circuit2: str
    reason: str


@dataclass
class Carryover:
    proposals: List[Proposal] = field(default_factory=list)
    lost: List[Lost] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)

    # ------------------------------------------------------------ reading
    def get(self, pid: str) -> Proposal:
        for p in self.proposals:
            if p.id == pid:
                return p
        raise KeyError(pid)

    def of_kind(self, kind: str) -> List[Proposal]:
        return [p for p in self.proposals if p.kind == kind]

    @property
    def undecided(self) -> List[Proposal]:
        return [p for p in self.proposals if not p.decided]

    def counts(self) -> Dict[str, int]:
        return {
            EXACT: len(self.of_kind(EXACT)),
            SALES_CODE: len(self.of_kind(SALES_CODE)),
            CHANGED: len(self.of_kind(CHANGED)),
            "kept": sum(1 for p in self.proposals if p.existing),
            "decided": sum(1 for p in self.proposals if p.decided),
            "undecided": len(self.undecided),
            "lost": len(self.lost),
            "written": sum(1 for p in self.proposals
                           if p.decision in (COPY, NEW) and p.result),
        }

    # ----------------------------------------------------------- deciding
    def decide(self, pid: str, decision: str, text: str = "") -> Proposal:
        """Record the engineer's decision for one row."""
        p = self.get(pid)
        if decision not in DECISION_LABEL:
            raise ValueError(f"unknown decision {decision!r}")
        if decision == NEW and not text.strip():
            raise ValueError("a new comment needs text")
        if decision == KEEP and not p.existing:
            raise ValueError("this row has no comment of its own to keep")
        p.decision = decision
        p.text = text.strip() if decision == NEW else ""
        return p

    def undo(self, pid: str) -> Proposal:
        """Back to undecided — never for an identical row, which needs no one."""
        p = self.get(pid)
        if p.kind != EXACT or p.existing:
            p.decision, p.text = None, ""
        return p

    def accept_suggestions(self, kind: Optional[str] = None) -> int:
        """Apply the suggestion to every undecided row that has one."""
        n = 0
        for p in self.undecided:
            if p.suggested and (kind is None or p.kind == kind):
                p.decision = p.suggested
                n += 1
        return n

    def decide_all(self, kind: str, decision: str) -> int:
        """One decision for every undecided row of a kind. NEW needs text per
        row, so it cannot be applied in bulk."""
        if decision == NEW:
            raise ValueError("a new comment is written row by row")
        n = 0
        for p in self.undecided:
            if p.kind == kind and (decision != KEEP or p.existing):
                p.decision = decision
                n += 1
        return n


def _signature(row: Row, keys: List[str]) -> tuple:
    return tuple(row.values.get(k, "") for k in keys)


def _same_wire(old: Row, new: Row) -> bool:
    """Same pin, and the same circuit on at least one side."""
    if old.pin != new.pin:
        return False
    return bool((old.circuit1 and old.circuit1 == new.circuit1)
                or (old.circuit2 and old.circuit2 == new.circuit2))


def _diffs(old: Row, new: Row, keys: List[str]) -> List[Diff]:
    return [Diff(k, old.values.get(k, ""), new.values.get(k, "")) for k in keys
            if old.values.get(k, "") != new.values.get(k, "")]


def match(old: Report, new: Report) -> Carryover:
    """Pair every commented old row with at most one new row."""
    out = Carryover()

    for name, new_sheet in new.sheets.items():
        old_sheet = old.sheets.get(name)
        if old_sheet is None:
            continue
        keys: List[str] = []
        for column in old_sheet.columns + new_sheet.columns:
            if column.key not in keys:
                keys.append(column.key)
        if {c.key for c in old_sheet.columns} != {c.key for c in new_sheet.columns}:
            out.notes.append(f"{name}: the column set changed between the "
                             "reports; a column present in only one is read as "
                             "blank in the other.")

        sources = [r for r in old_sheet.rows if r.comment]
        used: set = set()
        chosen: Dict[int, Tuple[Row, List[Diff], List[str]]] = {}

        # 1 · identical rows, in order, one old row each
        pool: Dict[tuple, List[Row]] = {}
        for src in sources:
            pool.setdefault(_signature(src, keys), []).append(src)
        for target in new_sheet.rows:
            queue = pool.get(_signature(target, keys))
            if queue:
                src = queue.pop(0)
                used.add(src.row)
                chosen[target.row] = (src, [], [])

        # 2 · the same wire, fewest differing columns first
        pairs = []
        for target in new_sheet.rows:
            if target.row in chosen:
                continue
            for src in sources:
                if src.row in used or not _same_wire(src, target):
                    continue
                diffs = _diffs(src, target, keys)
                pairs.append((len(diffs), target.row, src.row, target, src, diffs))
        pairs.sort(key=lambda p: (p[0], p[1], p[2]))
        for score, t_row, s_row, target, src, diffs in pairs:
            if t_row in chosen or s_row in used:
                continue
            tied = [p[4].comment for p in pairs
                    if p[1] == t_row and p[0] == score and p[2] != s_row
                    and p[4].comment != src.comment]
            chosen[t_row] = (src, diffs, sorted(set(tied)))
            used.add(s_row)

        # proposals, in the new report's row order
        for target in new_sheet.rows:
            if target.row not in chosen:
                continue
            src, diffs, alternatives = chosen[target.row]
            if target.comment and target.comment == src.comment:
                continue                      # already says what it said
            if not diffs:
                kind = EXACT
            elif all(is_sales_code(d.key) for d in diffs):
                kind = SALES_CODE
            else:
                kind = CHANGED
            if target.comment:
                suggested = decision = KEEP
            elif kind == EXACT:
                suggested = decision = COPY
            elif kind == SALES_CODE:
                suggested, decision = COPY, None
            else:
                suggested = decision = None
            out.proposals.append(Proposal(
                sheet=name, new_row=target.row, old_row=src.row,
                new_col=new_sheet.comment_col, old_col=old_sheet.comment_col,
                kind=kind, diffs=diffs, old_comment=src.comment,
                existing=target.comment, pin=target.pin,
                circuit1=target.circuit1, circuit2=target.circuit2,
                suggested=suggested, decision=decision,
                alternatives=alternatives))

        for src in sources:
            if src.row not in used:
                out.lost.append(Lost(name, src.row, src.comment, src.pin,
                                     src.circuit1, src.circuit2,
                                     "no row in the new report has this pin "
                                     "and circuit"))

    for name, old_sheet in old.sheets.items():
        if name in new.sheets:
            continue
        commented = [r for r in old_sheet.rows if r.comment]
        if commented:
            out.notes.append(f"{name}: in the old report but not the new one — "
                             f"its {len(commented)} comment(s) did not carry.")
        for src in commented:
            out.lost.append(Lost(name, src.row, src.comment, src.pin,
                                 src.circuit1, src.circuit2,
                                 "this inline sheet is not in the new report"))
    return out


# ------------------------------------------------------------------ writing
def apply(old_data: bytes, new_data: bytes, carryover: Carryover,
          keep_vba: bool = False) -> bytes:
    """The new report with the decided comments written — and nothing else
    touched. Refuses while any row is undecided: that is the review gate."""
    if carryover.undecided:
        raise UndecidedRows(len(carryover.undecided))
    old_wb = load_workbook(io.BytesIO(old_data))
    new_wb = load_workbook(io.BytesIO(new_data), keep_vba=keep_vba)
    for p in carryover.proposals:
        cell = new_wb[p.sheet].cell(p.new_row, p.new_col)
        if p.decision == COPY:
            source = old_wb[p.sheet].cell(p.old_row, p.old_col)
            cell.value = source.value
            cell.fill = copy(source.fill)
            cell.font = copy(source.font)
        elif p.decision == NEW:
            cell.value = p.text
        elif p.decision == BLANK:
            cell.value = None
        # KEEP: the new report's own comment stays exactly as it is
    buffer = io.BytesIO()
    new_wb.save(buffer)
    return buffer.getvalue()


def output_name(new_name: str) -> str:
    """``IP_Inline_Report.xlsx`` → ``IP_Inline_Report_commented.xlsx``."""
    stem, dot, ext = (new_name or "Inline_Report.xlsx").rpartition(".")
    if not dot:
        return f"{new_name}_commented.xlsx"
    return f"{stem}_commented.{ext}"
