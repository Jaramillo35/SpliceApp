"""The sales-code band of a master-complexity family worksheet.

Every family worksheet carries a header band on **row 6** whose cells name what
the columns below them hold: ``CPOS Packages``, ``Markets``, ``Optional
Features``, ``Standard Features``, ``Requested Field for EBOM``,
``RELEASE/EBOM STRING`` and so on. The sales codes sit on **row 9** between two
of those anchors:

    Optional Features  ...........................  RELEASE/EBOM STRING
    ^ first code column                            ^ first non-code column

Left of the band are package columns (``PC3``, ``PC5 AWD``) and, directly
before ``Optional Features`` under a ``Markets`` header, the market columns
(``YAA``, ``YAC``). Neither is in the band. Packages are never sales codes;
markets are read separately and offered — the engineer decides, per family,
whether a market code becomes a column of the individual file, and nothing is
added unasked. Everything from the right anchor onward is release and analyst
bookkeeping. Packages look like sales codes to a regex — ``PC3`` and ``AWD``
are three characters — and that is why the workbench used to need a DTx export
to say which row-9 tokens were real. The band makes the DTx unnecessary: the
master already says where the codes are.

What made this fragile before it was written down, in the reference master:

* the left anchor is ``Optional Features`` on most sheets and
  ``Optional\\nFeatures`` — with a line break — on three of them. A plain
  substring test misses those three, and they were exactly the harnesses whose
  codes were not detected;
* one sheet (a jumper) has its codes on **row 8**, with row 9 empty in the
  band. Row 9 is the rule; row 8 is a reported fallback, never a silent one;
* one sheet (a seat) has no ``RELEASE/EBOM STRING`` cell at all — the row-6
  header is blank where it should be. On every sheet the columns after the
  codes are EBOM bookkeeping (``Requested Field for EBOM``, ``EBOM Analyst`` on
  the row above, ``EBOM X-Check``), so the band ends at the first header after
  the left anchor that mentions EBOM. That is the specified anchor's column or
  one left of it on the sheets that have the anchor, and still an answer on
  the one that does not. Which text closed the band is recorded;
* cells carry prose beside the codes — ``JRK (Includes JPG)``, ``XAK (NO
  XAC)`` — so tokens are matched on word boundaries: ``JRK`` and ``JPG``,
  never ``INC`` and ``LUD`` out of ``Includes``;
* the variant marker columns (``LEFT`` / ``RIGHT``) sit inside the band on
  partitioned sheets; they never match the token shape and fall out of their
  own accord. ``CUP`` / ``CM5`` / ``CVM`` do match, and are kept: a console
  variant marker is also that variant's sales code.

The output per cell is the shape asked for: column, the raw expression, the
codes in it, and the operators in it. The operators are recorded, not
interpreted — ``+`` is AND, ``/`` is OR, ``-`` is NOT, ``=`` is a package
equivalence and ``()`` groups, and the rest of the workbench decides what to
do with an expression that uses them.
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from typing import Dict, Iterable, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

SALES_CODE_ROW = 9
FALLBACK_ROW = 8
FEATURE_ROW = 7
#: the anchors are on row 6 in every master seen; searched a little wider so
#: a master that gains a title row does not become undetectable
ANCHOR_ROWS = range(1, 13)

LEFT_ANCHOR = "optional features"
#: the market columns sit under this header, directly before the left
#: anchor; read separately and offered, never added to the band
MARKETS_ANCHOR = "markets"
RIGHT_ANCHOR = "release/ebom string"
#: any header mentioning this closes the band — the specified anchor and the
#: EBOM bookkeeping columns that always precede or replace it
RIGHT_FAMILY = "ebom"

#: The characters that are logic, not codes. ``%`` has not been seen in a
#: master yet; it is in the specification, so it is recognised.
OPERATORS = "+-/=%()"

#: A sales code is three capitals/digits standing alone — on word boundaries,
#: so prose beside a code (``Includes``, ``NAFTA``) contributes nothing, and
#: case-sensitive as specified, so ``ckt`` in ``No E12 ckt`` is not a code
#: either. Masters write codes in capitals; a lower-case token is prose.
TOKEN = re.compile(r"\b[A-Z0-9]{3}\b")

def normalize(text: object) -> str:
    """Whitespace collapsed to single spaces, lower-cased — for matching."""
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def parse_cell(text: object) -> Tuple[List[str], List[str]]:
    """``(sales_codes, operators)`` found in one row-9 cell, each in order.

    Codes are unique and keep first-seen order; operators keep every
    occurrence, because ``A/B/C`` has two ORs and that is part of the shape.
    """
    raw = str(text or "")
    codes: List[str] = []
    for token in TOKEN.findall(raw):
        if token not in codes:
            codes.append(token)
    operators = [ch for ch in raw if ch in OPERATORS]
    return codes, operators


@dataclass
class SalesCodeCell:
    """One row-9 cell inside the band, read as the specification asks."""

    column: int                       # 1-based worksheet column
    column_name: str                  # its letter, e.g. "AU"
    raw_expression: str
    sales_codes: List[str] = field(default_factory=list)
    operators: List[str] = field(default_factory=list)
    feature: str = ""                 # row 7, the description above the code
    row: int = SALES_CODE_ROW

    def as_dict(self) -> dict:
        return {"column_name": self.column_name,
                "raw_expression": self.raw_expression,
                "sales_codes": list(self.sales_codes),
                "operators": list(self.operators)}

    @property
    def is_single(self) -> bool:
        return len(self.sales_codes) == 1 and not self.operators

    @property
    def is_or_list(self) -> bool:
        """Only ``/`` between codes — safe to split into independent columns."""
        return (len(self.sales_codes) > 1
                and bool(self.operators)
                and all(op == "/" for op in self.operators))

    @property
    def is_equality(self) -> bool:
        return "=" in self.operators and all(op == "=" for op in self.operators)

    @property
    def is_combined(self) -> bool:
        """AND, NOT or grouping — cannot be split without changing its meaning."""
        return any(op in "+-%()" for op in self.operators)


@dataclass
class Band:
    worksheet: str
    anchor_row: int
    start_col: int                    # the left anchor's column, inclusive
    end_col: int                      # the right anchor's column, exclusive
    left_anchor: str = ""             # the header text that opened the band
    right_anchor: str = ""            # the header text that closed the band
    row_used: int = SALES_CODE_ROW
    cells: List[SalesCodeCell] = field(default_factory=list)
    #: The market columns (YAA, YAC) under the ``Markets`` header before the
    #: band. Not codes of the band; offered to the engineer as optional
    #: columns of the individual file.
    markets: List[SalesCodeCell] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)

    @property
    def market_codes(self) -> List[str]:
        out: List[str] = []
        for cell in self.markets:
            for code in cell.sales_codes:
                if code not in out:
                    out.append(code)
        return out

    @property
    def codes(self) -> List[str]:
        """Every code in the band, unique, in column order."""
        out: List[str] = []
        for cell in self.cells:
            for code in cell.sales_codes:
                if code not in out:
                    out.append(code)
        return out

    @property
    def columns(self) -> range:
        return range(self.start_col, self.end_col)

    def as_dicts(self) -> List[dict]:
        return [cell.as_dict() for cell in self.cells]


# ------------------------------------------------------------- worksheets
def _row(ws, index: int) -> List[object]:
    """One row's values, on a read-only or a normal worksheet alike."""
    try:
        return list(next(ws.iter_rows(min_row=index, max_row=index,
                                      values_only=True)))
    except StopIteration:
        return []


def find_anchors(ws) -> Optional[Tuple[int, int, int, Optional[int], str, str]]:
    """``(row, left column, right column, markets column, left text, right
    text)``, or None.

    The left anchor is ``Optional Features``. The ``Markets`` header that
    precedes it on the same row, when there is one, is returned beside it so
    the market columns can be read separately. The right boundary is the first
    header after it — on the same row or the one above — that mentions EBOM:
    ``RELEASE/EBOM STRING`` as specified, or the ``Requested Field for EBOM``
    / ``EBOM Analyst`` bookkeeping that always sits there, on the sheets whose
    ``RELEASE/EBOM STRING`` cell is missing. An EBOM header alone (the summary
    sheets have several) is not a band: the left anchor is required.
    """
    for r in ANCHOR_ROWS:
        values = _row(ws, r)
        optional = next((c for c, v in enumerate(values, start=1)
                         if LEFT_ANCHOR in normalize(v)), None)
        if optional is None:
            continue
        markets = next((c for c in range(optional - 1, 0, -1)
                        if MARKETS_ANCHOR in normalize(values[c - 1])), None)
        left_text = str(values[optional - 1]).strip()
        above = _row(ws, r - 1) if r > 1 else []
        for c in range(optional + 1, max(len(values), len(above)) + 1):
            for row in (values, above):
                text = normalize(row[c - 1]) if c - 1 < len(row) else ""
                if RIGHT_FAMILY in text:
                    return (r, optional, c, markets, left_text,
                            str(row[c - 1]).strip())
        return None
    return None


def is_family_sheet(ws) -> bool:
    return find_anchors(ws) is not None


def read_band(ws, worksheet: str = "") -> Optional[Band]:
    """The band of one worksheet, or None when it has no anchors."""
    anchors = find_anchors(ws)
    if anchors is None:
        return None
    anchor_row, start, end, markets_col, left_text, right_text = anchors
    band = Band(worksheet=worksheet or getattr(ws, "title", ""),
                anchor_row=anchor_row, start_col=start, end_col=end,
                left_anchor=left_text, right_anchor=right_text)
    features = _row(ws, FEATURE_ROW)
    if markets_col is not None:
        band.markets = _cells(_row(ws, SALES_CODE_ROW), features,
                              markets_col, start, SALES_CODE_ROW)
    # The specified anchor is usually one column RIGHT of the header that
    # closes the band ('Requested Field for EBOM' precedes it on 29 of 32
    # sheets). Only its absence from the whole row is worth a note.
    row_text = _row(ws, anchor_row)
    if not any(RIGHT_ANCHOR in normalize(v) for v in row_text[start:]):
        band.notes.append(
            f"No 'RELEASE/EBOM STRING' header on row {anchor_row}; the band was "
            f"closed at column {get_column_letter(end)} by {right_text!r}.")

    cells = _cells(_row(ws, SALES_CODE_ROW), features, start, end, SALES_CODE_ROW)
    if not cells:
        # Row 9 is the rule. When it is empty across the whole band the codes
        # are usually one row up — one jumper sheet in the reference master
        # does this — and reading nothing would report a family with no sales
        # codes at all. So the fallback is taken and SAID, never silent.
        fallback = _cells(_row(ws, FALLBACK_ROW), features, start, end, FALLBACK_ROW)
        if fallback:
            band.row_used = FALLBACK_ROW
            band.notes.append(
                f"Row {SALES_CODE_ROW} is empty between the anchors; the "
                f"{len(fallback)} code cell(s) were read from row {FALLBACK_ROW}.")
            cells = fallback
        else:
            band.notes.append(
                f"No sales codes on row {SALES_CODE_ROW} (or {FALLBACK_ROW}) "
                f"between {get_column_letter(start)} and {get_column_letter(end)}.")
    band.cells = cells
    return band


def _cells(values: List[object], features: List[object], start: int, end: int,
           row: int) -> List[SalesCodeCell]:
    out: List[SalesCodeCell] = []
    for c in range(start, end):
        value = values[c - 1] if c - 1 < len(values) else None
        raw = str(value).strip() if value is not None else ""
        if not raw:
            continue
        # LEFT / RIGHT / DRIVER / PASSENGER never match the token shape, so
        # they fall out here. CUP, CM5 and CVM DO match — and they are sales
        # codes as well as the console variant they name, so they are kept;
        # which variant a part belongs to is partition.py's question.
        codes, operators = parse_cell(raw)
        if not codes:
            continue                  # prose, or a sub-header like "Standard Features"
        feature = features[c - 1] if c - 1 < len(features) else None
        out.append(SalesCodeCell(
            column=c, column_name=get_column_letter(c), raw_expression=raw,
            sales_codes=codes, operators=operators,
            feature=str(feature).strip() if feature is not None else "", row=row))
    return out


# --------------------------------------------------------------- workbooks
def read_master(master_bytes: bytes,
                worksheets: Optional[Iterable[str]] = None) -> Dict[str, Band]:
    """Every family worksheet's band, keyed by sheet name.

    A sheet without both anchors is not a family sheet (the change log, the
    summary, the markup guide) and is left out rather than reported empty.
    ``worksheets`` restricts the result to the names given.
    """
    from splice.common.errors import SpliceError

    if not master_bytes:
        raise SpliceError("Master complexity workbook is empty or was not provided.")
    try:
        wb = load_workbook(io.BytesIO(master_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise SpliceError(f"Could not read the master complexity workbook: {exc}") from exc

    wanted = set(worksheets) if worksheets is not None else None
    out: Dict[str, Band] = {}
    try:
        for name in wb.sheetnames:
            if wanted is not None and name not in wanted:
                continue
            band = read_band(wb[name], name)
            if band is not None:
                out[name] = band
    finally:
        wb.close()
    logger.info("Master band: %d family worksheet(s), %d code(s).",
                len(out), sum(len(b.codes) for b in out.values()))
    return out


def family_worksheets(master_bytes: bytes) -> List[str]:
    """The sheets that are harness families: a band with at least one code.

    A master's TEMPLATE SHEET carries the anchors and nothing between them;
    ``read_master`` still returns it, with its note, but it is not a family
    an engineer can open.
    """
    return [name for name, b in read_master(master_bytes).items() if b.codes]
